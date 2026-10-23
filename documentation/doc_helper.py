import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_base_document(header_title="PANADERÍA SVETLANA — SISTEMA INTEGRAL DE GESTIÓN"):
    doc = Document()
    
    # Configure page margins to 1 inch (2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Configure Header
        header = section.header
        p_head = header.paragraphs[0]
        p_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_head = p_head.add_run(header_title)
        r_head.font.name = "Arial"
        r_head.font.size = Pt(8.5)
        r_head.font.color.rgb = RGBColor(148, 163, 184) # Slate-400
        
        # Configure Footer
        footer = section.footer
        p_foot = footer.paragraphs[0]
        p_foot.alignment = WD_ALIGN_PARAGRAPH.LEFT
        
        r_foot_left = p_foot.add_run("Panadería Svetlana | Documentación Oficial del Sistema                    ")
        r_foot_left.font.name = "Arial"
        r_foot_left.font.size = Pt(8.5)
        r_foot_left.font.color.rgb = RGBColor(148, 163, 184)
        
        r_foot_right = p_foot.add_run("Página ")
        r_foot_right.font.name = "Arial"
        r_foot_right.font.size = Pt(8.5)
        r_foot_right.font.color.rgb = RGBColor(148, 163, 184)
        
        fld_page = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>')
        p_foot._p.append(fld_page)
        
        r_foot_mid = p_foot.add_run(" de ")
        r_foot_mid.font.name = "Arial"
        r_foot_mid.font.size = Pt(8.5)
        r_foot_mid.font.color.rgb = RGBColor(148, 163, 184)
        
        fld_numpages = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="NUMPAGES"/>')
        p_foot._p.append(fld_numpages)
        
    # Configure Normal Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Arial'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(30, 41, 59) # Slate-800
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)
    
    return doc

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_cell_borders(cell, top="none", bottom="none", left="none", right="none", 
                     color="CBD5E1", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    def bdr_xml(side, val):
        if val == "none":
            return f'<w:{side} {nsdecls("w")} w:val="none"/>'
        return f'<w:{side} {nsdecls("w")} w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
    borders_xml = f'''
        <w:tcBorders {nsdecls("w")}>
            {bdr_xml("top", top)}
            {bdr_xml("left", left)}
            {bdr_xml("bottom", bottom)}
            {bdr_xml("right", right)}
        </w:tcBorders>
    '''
    tcPr.append(parse_xml(borders_xml))

def add_main_cover(doc, main_title, subtitle, doc_code, version, date_str, classification, audience, metadata_rows=None):
    # Top spacing
    p_top = doc.add_paragraph()
    p_top.paragraph_format.space_before = Pt(36)
    p_top.paragraph_format.space_after = Pt(8)
    
    # Pre-title
    r_pre = p_top.add_run("PANADERÍA SVETLANA — SISTEMA WEB Y GESTIÓN DE PRODUCCIÓN")
    r_pre.font.name = "Arial"
    r_pre.font.size = Pt(11)
    r_pre.font.bold = True
    r_pre.font.color.rgb = RGBColor(37, 99, 235) # Blue-600
    
    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(12)
    r_title = p_title.add_run(main_title)
    r_title.font.name = "Arial"
    r_title.font.size = Pt(26)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(30, 58, 138) # Navy #1E3A8A
    
    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(24)
    r_sub = p_sub.add_run(subtitle)
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(15)
    r_sub.font.color.rgb = RGBColor(71, 85, 105) # Slate-600
    
    # Badge / Document info table
    table = doc.add_table(rows=6, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    col_widths = [Inches(2.2), Inches(4.3)]
    meta = [
        ("Código de Documento:", doc_code),
        ("Versión Oficial:", version),
        ("Fecha de Emisión / Revisión:", date_str),
        ("Clasificación de Seguridad:", classification),
        ("Audiencia Objetivo:", audience),
        ("Alcance Operativo:", "Sucursal Matriz (Panadería y Producción Central) y Sucursal Secundaria")
    ]
    if metadata_rows:
        meta.extend(metadata_rows)
        while len(table.rows) < len(meta):
            table.add_row()
            
    for i, (label, val) in enumerate(meta):
        row = table.rows[i]
        
        # Label cell
        c0 = row.cells[0]
        c0.width = col_widths[0]
        set_cell_shading(c0, "1E3A8A")
        set_cell_margins(c0, 100, 100, 140, 140)
        set_cell_borders(c0, top="single", bottom="single", left="single", right="single", color="CBD5E1", sz="4")
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        p0.paragraph_format.line_spacing = 1.0
        r0 = p0.add_run(label)
        r0.font.name = "Arial"
        r0.font.size = Pt(10)
        r0.font.bold = True
        r0.font.color.rgb = RGBColor(255, 255, 255)
        
        # Value cell
        c1 = row.cells[1]
        c1.width = col_widths[1]
        set_cell_shading(c1, "F8FAFC" if i % 2 == 0 else "FFFFFF")
        set_cell_margins(c1, 100, 100, 140, 140)
        set_cell_borders(c1, top="single", bottom="single", left="single", right="single", color="CBD5E1", sz="4")
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        p1.paragraph_format.line_spacing = 1.0
        r1 = p1.add_run(val)
        r1.font.name = "Arial"
        r1.font.size = Pt(10)
        r1.font.color.rgb = RGBColor(30, 41, 59)
        
    p_bottom = doc.add_paragraph()
    p_bottom.paragraph_format.space_before = Pt(36)
    p_bottom.paragraph_format.space_after = Pt(0)
    r_bot = p_bottom.add_run("CONTROL DE CONFIDENCIALIDAD: Este documento contiene especificaciones técnicas, operativas y de procesos propiedad exclusiva de Panadería Svetlana. Prohibida su reproducción, distribución o divulgación total o parcial sin autorización expresa por escrito de la Dirección General.")
    r_bot.font.name = "Arial"
    r_bot.font.size = Pt(9)
    r_bot.font.italic = True
    r_bot.font.color.rgb = RGBColor(100, 116, 139)
    
    doc.add_page_break()

def add_divider_cover(doc, part_number, part_title, part_subtitle, description, topics_list):
    doc.add_page_break()
    
    # Spacing
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(60)
    
    # Part Number Badge
    p_part = doc.add_paragraph()
    p_part.paragraph_format.space_after = Pt(6)
    r_part = p_part.add_run(f"SUBDIVISIÓN {part_number}")
    r_part.font.name = "Arial"
    r_part.font.size = Pt(13)
    r_part.font.bold = True
    r_part.font.color.rgb = RGBColor(37, 99, 235) # Blue-600
    
    # Part Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(8)
    r_title = p_title.add_run(part_title)
    r_title.font.name = "Arial"
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(30, 58, 138) # Navy
    
    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(18)
    r_sub = p_sub.add_run(part_subtitle)
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(14)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(71, 85, 105)
    
    # Scope box
    add_callout(doc, description, title="ALCANCE Y PROPÓSITO DE ESTA SUBDIVISIÓN:", callout_type="info")
    
    # Modules breakdown table
    p_mod_title = doc.add_paragraph()
    p_mod_title.paragraph_format.space_before = Pt(16)
    p_mod_title.paragraph_format.space_after = Pt(6)
    r_mod_title = p_mod_title.add_run("Contenido y Módulos Desarrollados en esta Sección:")
    r_mod_title.font.name = "Arial"
    r_mod_title.font.size = Pt(11)
    r_mod_title.font.bold = True
    r_mod_title.font.color.rgb = RGBColor(30, 41, 59)
    
    headers = ["Módulo / Capítulo", "Objetivo Operativo / Técnico", "Audiencia Responsable"]
    add_styled_table(doc, headers, topics_list, col_widths=[Inches(1.8), Inches(3.2), Inches(1.5)])
    
    doc.add_page_break()

def add_h1(doc, text, page_break_before=False):
    if page_break_before:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(17)
    r.font.bold = True
    r.font.color.rgb = RGBColor(30, 58, 138) # Navy #1E3A8A
    return p

def add_h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(37, 99, 235) # Blue #2563EB
    return p

def add_h3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(12.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(51, 65, 85) # Slate-700
    return p

def add_p(doc, text, bold_prefix=None, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + " ")
        r_pre.font.name = "Arial"
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(30, 41, 59)
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(12)
    r.font.italic = italic
    r.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_bullet(doc, text, bold_prefix=None, level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + " ")
        r_pre.font.name = "Arial"
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(30, 41, 59)
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_callout(doc, text, title="NOTA:", callout_type="info"):
    colors = {
        "info": {"border": "2563EB", "fill": "EFF6FF", "title_rgb": RGBColor(37, 99, 235)},
        "warning": {"border": "D97706", "fill": "FFFBEB", "title_rgb": RGBColor(217, 119, 6)},
        "success": {"border": "059669", "fill": "ECFDF5", "title_rgb": RGBColor(5, 150, 105)},
        "rule": {"border": "1E3A8A", "fill": "F8FAFC", "title_rgb": RGBColor(30, 58, 138)},
    }
    cfg = colors.get(callout_type, colors["info"])
    
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.5)
    set_cell_shading(cell, cfg["fill"])
    set_cell_margins(cell, 120, 120, 160, 160)
    set_cell_borders(cell, top="none", bottom="none", left="single", right="none", 
                     color=cfg["border"], sz="24") # sz 24 = 3pt
    
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(2)
    
    r_tit = p.add_run(title + " ")
    r_tit.font.name = "Arial"
    r_tit.font.size = Pt(11)
    r_tit.font.bold = True
    r_tit.font.color.rgb = cfg["title_rgb"]
    
    r_body = p.add_run(text)
    r_body.font.name = "Arial"
    r_body.font.size = Pt(11)
    r_body.font.color.rgb = RGBColor(30, 41, 59)
    
    # spacing after table
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(4)

def add_styled_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Header Row
    header_row = table.rows[0]
    for j, h_text in enumerate(headers):
        cell = header_row.cells[j]
        if col_widths and j < len(col_widths):
            cell.width = col_widths[j]
        set_cell_shading(cell, "1E3A8A") # Navy
        set_cell_margins(cell, 120, 120, 140, 140)
        set_cell_borders(cell, top="single", bottom="single", left="single", right="single", color="94A3B8", sz="4")
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(h_text)
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    # Data Rows
    for i, row_data in enumerate(rows):
        row = table.rows[i + 1]
        fill_hex = "F8FAFC" if i % 2 == 1 else "FFFFFF"
        for j, cell_text in enumerate(row_data):
            cell = row.cells[j]
            if col_widths and j < len(col_widths):
                cell.width = col_widths[j]
            set_cell_shading(cell, fill_hex)
            set_cell_margins(cell, 100, 100, 140, 140)
            set_cell_borders(cell, top="single", bottom="single", left="single", right="single", color="CBD5E1", sz="4")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            r = p.add_run(str(cell_text))
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(30, 41, 59)
            
    # spacing after table
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(6)

def add_image_with_caption(doc, image_path, caption_text, width_inches=6.2):
    if not os.path.exists(image_path):
        add_callout(doc, f"Imagen no encontrada en la ruta: {image_path}", title="ERROR DE RECURSO:", callout_type="warning")
        return
        
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(4)
    run_img = p_img.add_run()
    run_img.add_picture(image_path, width=Inches(width_inches))
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(0)
    p_cap.paragraph_format.space_after = Pt(10)
    r_cap = p_cap.add_run(caption_text)
    r_cap.font.name = "Arial"
    r_cap.font.size = Pt(9.5)
    r_cap.font.italic = True
    r_cap.font.color.rgb = RGBColor(100, 116, 139) # Slate-500

def add_screenshot_placeholder(doc, title, url_route, role_responsible, key_elements, caption_text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.5)
    set_cell_shading(cell, "F8FAFC")
    set_cell_margins(cell, 160, 160, 180, 180)
    set_cell_borders(cell, top="single", bottom="single", left="single", right="single", color="94A3B8", sz="6")
    
    p0 = cell.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_after = Pt(6)
    r0 = p0.add_run(f"[ESPACIO RESERVADO PARA CAPTURA DE PANTALLA: {title.upper()}]")
    r0.font.name = "Arial"
    r0.font.size = Pt(11)
    r0.font.bold = True
    r0.font.color.rgb = RGBColor(30, 58, 138)
    
    p1 = cell.add_paragraph()
    p1.paragraph_format.line_spacing = 1.15
    p1.paragraph_format.space_after = Pt(3)
    r1_pre = p1.add_run("Ruta en Sistema / URL: ")
    r1_pre.font.name = "Arial"
    r1_pre.font.size = Pt(10)
    r1_pre.font.bold = True
    r1_val = p1.add_run(url_route)
    r1_val.font.name = "Arial"
    r1_val.font.size = Pt(10)
    r1_val.font.color.rgb = RGBColor(37, 99, 235)
    
    p2 = cell.add_paragraph()
    p2.paragraph_format.line_spacing = 1.15
    p2.paragraph_format.space_after = Pt(3)
    r2_pre = p2.add_run("Rol Autorizado para la Captura: ")
    r2_pre.font.name = "Arial"
    r2_pre.font.size = Pt(10)
    r2_pre.font.bold = True
    r2_val = p2.add_run(role_responsible)
    r2_val.font.name = "Arial"
    r2_val.font.size = Pt(10)
    
    p3 = cell.add_paragraph()
    p3.paragraph_format.line_spacing = 1.15
    p3.paragraph_format.space_after = Pt(3)
    r3_pre = p3.add_run("Elementos Requeridos en la Captura: ")
    r3_pre.font.name = "Arial"
    r3_pre.font.size = Pt(10)
    r3_pre.font.bold = True
    r3_val = p3.add_run(key_elements)
    r3_val.font.name = "Arial"
    r3_val.font.size = Pt(10)
    
    p4 = cell.add_paragraph()
    p4.paragraph_format.line_spacing = 1.15
    p4.paragraph_format.space_after = Pt(0)
    r4_pre = p4.add_run("Resolución / Formato Recomendado: ")
    r4_pre.font.name = "Arial"
    r4_pre.font.size = Pt(10)
    r4_pre.font.bold = True
    r4_val = p4.add_run("1440x900 px o 1920x1080 px (PNG sin compresión). Insertar imagen reemplazando este recuadro.")
    r4_val.font.name = "Arial"
    r4_val.font.size = Pt(9.5)
    r4_val.font.italic = True
    r4_val.font.color.rgb = RGBColor(100, 116, 139)
    
    # Caption below table
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(10)
    r_cap = p_cap.add_run(caption_text)
    r_cap.font.name = "Arial"
    r_cap.font.size = Pt(9.5)
    r_cap.font.italic = True
    r_cap.font.color.rgb = RGBColor(100, 116, 139)

def add_code_block(doc, code_str):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.5)
    set_cell_shading(cell, "F1F5F9") # Slate-100
    set_cell_margins(cell, 120, 120, 140, 140)
    set_cell_borders(cell, top="single", bottom="single", left="single", right="single", color="CBD5E1", sz="4")
    
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(code_str)
    r.font.name = "Consolas"
    r.font.size = Pt(9.0)
    r.font.color.rgb = RGBColor(30, 41, 59)
    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(6)
